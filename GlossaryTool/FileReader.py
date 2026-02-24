import os
import zipfile
import json
import warnings

# Suppress warnings from libraries if necessary
warnings.filterwarnings("ignore")

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import ebooklib
    from ebooklib import epub
except ImportError:
    ebooklib = None
    epub = None

try:
    from pdfminer.high_level import extract_text as pdf_extract_text
except ImportError:
    pdf_extract_text = None

class FileReader:
    @staticmethod
    def read_file(file_path: str) -> str:
        """
        Reads content from a file based on its extension.
        Returns the text content as a string.
        """
        if not os.path.exists(file_path):
            return ""

        ext = os.path.splitext(file_path)[1].lower()
        
        try:
            if ext in ['.txt', '.md', '.xml', '.html', '.log', '.lrc', '.srt', '.ass', '.vtt']:
                return FileReader._read_txt(file_path)
            elif ext == '.json':
                return FileReader._read_json(file_path)
            elif ext == '.docx':
                return FileReader._read_docx(file_path)
            elif ext == '.xlsx':
                return FileReader._read_xlsx(file_path)
            elif ext == '.epub':
                return FileReader._read_epub(file_path)
            elif ext == '.pdf':
                return FileReader._read_pdf(file_path)
            else:
                # Try reading as text for unknown extensions
                return FileReader._read_txt(file_path)
        except Exception as e:
            print(f"Error reading file {file_path}: {e}")
            return ""

    @staticmethod
    def _read_txt(file_path):
        try:
            # Try utf-8 first
            with open(file_path, 'r', encoding='utf-8') as f:
                return f.read()
        except UnicodeDecodeError:
            try:
                # Try with errors ignored or other encodings if needed
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
            except Exception:
                return ""

    @staticmethod
    def _read_json(file_path):
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                data = json.load(f)
            # Recursively extract strings from JSON
            return FileReader._extract_json_strings(data)
        except Exception:
            return FileReader._read_txt(file_path) # Fallback to raw text

    @staticmethod
    def _extract_json_strings(data):
        texts = []
        if isinstance(data, dict):
            # 针对翻译数据的特殊处理：如果存在 'src' 字段，仅提取它
            if 'src' in data and isinstance(data['src'], str):
                texts.append(data['src'])
            else:
                for v in data.values():
                    result = FileReader._extract_json_strings(v)
                    if result:
                        texts.append(result)
        elif isinstance(data, list):
            for item in data:
                result = FileReader._extract_json_strings(item)
                if result:
                    texts.append(result)
        elif isinstance(data, str):
            texts.append(data)
        return "\n".join(texts)

    @staticmethod
    def _read_docx(file_path):
        if not docx: 
            print("python-docx not installed")
            return ""
        try:
            doc = docx.Document(file_path)
            return "\n".join([para.text for para in doc.paragraphs])
        except Exception as e:
            print(f"docx read error: {e}")
            return ""

    @staticmethod
    def _read_xlsx(file_path):
        if not openpyxl: 
            print("openpyxl not installed")
            return ""
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            text = []
            for sheet in wb:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " ".join([str(cell) for cell in row if cell is not None])
                    text.append(row_text)
            return "\n".join(text)
        except Exception as e:
            print(f"xlsx read error: {e}")
            return ""

    @staticmethod
    def _read_epub(file_path):
        if not epub or not BeautifulSoup: 
            print("ebooklib or beautifulsoup4 not installed")
            return ""
        try:
            book = epub.read_epub(file_path)
            text = []
            # Iterate through all items
            for item in book.get_items():
                if item.get_type() == ebooklib.ITEM_DOCUMENT:
                    soup = BeautifulSoup(item.get_content(), 'html.parser')
                    text.append(soup.get_text())
            return "\n".join(text)
        except Exception as e:
            print(f"epub read error: {e}")
            return ""

    @staticmethod
    def _read_pdf(file_path):
        if not pdf_extract_text: 
            print("pdfminer.six not installed")
            return ""
        try:
            return pdf_extract_text(file_path)
        except Exception as e:
            print(f"pdf read error: {e}")
            return ""
