"""Versioned public-table bundles, with legacy-file detection."""

import json
import os
from pathlib import Path
import re
import tempfile

import openpyxl

from Frontend.services.editor import TableConfigHelper


class PublicTableFileHelper:
    FORMAT = "ainiee.public_tables"
    VERSION = 1
    METADATA_SHEET = "_AiNiee_Tables"
    METADATA_HEADERS = ("sheet", "name", "enabled", "match_options")
    INVALID = "表格文件格式无效或已损坏。"
    UNSUPPORTED_VERSION = "不支持的表格文件版本。"
    WRONG_KIND = "文件中的表格分类与当前分类不一致。"
    EMPTY = "表格文件中没有可导入的子表。"
    UNSUPPORTED_FILE = "不支持的文件类型。请选择 .json 或 .xlsx"
    TEXT_TOO_LONG = "表格文本过长，无法导出为 XLSX，请使用 JSON。"

    @classmethod
    def load_bundle(cls, path, kind, keys):
        """Return bundled tables, or None for an existing single-table format."""
        path = Path(path)
        extension = path.suffix.lower()
        if extension == ".json":
            try:
                with path.open("r", encoding="utf-8-sig") as reader:
                    payload = json.load(reader)
            except json.JSONDecodeError as error:
                raise ValueError(cls.INVALID) from error
            if not cls._looks_like_bundle(payload):
                return None
            return cls._validate_bundle(payload, kind, keys)
        if extension == ".xlsx":
            return cls._load_workbook(path, kind, keys)
        raise ValueError(cls.UNSUPPORTED_FILE)

    @classmethod
    def save_bundle(cls, path, kind, tables, keys, headers):
        """Write a complete bundle without replacing the destination on failure."""
        path = Path(path)
        extension = path.suffix.lower()
        if extension not in (".json", ".xlsx"):
            raise ValueError(cls.UNSUPPORTED_FILE)
        payload = {"format": cls.FORMAT, "version": cls.VERSION, "kind": kind, "tables": tables}
        payload["tables"] = cls._validate_bundle(payload, kind, keys)
        descriptor, temporary_name = tempfile.mkstemp(prefix=f".{path.stem}-", suffix=extension, dir=path.parent)
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            if extension == ".json":
                with temporary_path.open("w", encoding="utf-8") as writer:
                    json.dump(payload, writer, ensure_ascii=False, indent=2)
            else:
                cls._save_workbook(temporary_path, payload, keys, headers)
            os.replace(temporary_path, path)
        finally:
            temporary_path.unlink(missing_ok=True)

    @classmethod
    def _looks_like_bundle(cls, payload):
        if not isinstance(payload, dict):
            return False
        marker = payload.get("format")
        # A damaged envelope must not turn into glossary entries named "kind",
        # "version", or "tables" in the legacy key/value importer.
        return (
            isinstance(marker, str) and marker.startswith("ainiee.")
            or isinstance(payload.get("tables"), (list, dict))
        )

    @classmethod
    def _validate_bundle(cls, payload, kind, keys):
        if not isinstance(payload, dict) or payload.get("format") != cls.FORMAT:
            raise ValueError(cls.INVALID)
        if type(payload.get("version")) is not int or payload["version"] != cls.VERSION:
            raise ValueError(cls.UNSUPPORTED_VERSION)
        if kind not in TableConfigHelper.KINDS or payload.get("kind") not in TableConfigHelper.KINDS:
            raise ValueError(cls.INVALID)
        if payload["kind"] != kind:
            raise ValueError(cls.WRONG_KIND)
        tables = payload.get("tables")
        if not isinstance(tables, list):
            raise ValueError(cls.INVALID)
        if not tables:
            raise ValueError(cls.EMPTY)
        result = []
        for table in tables:
            if not isinstance(table, dict):
                raise ValueError(cls.INVALID)
            name, rows = table.get("name"), table.get("rows")
            if not isinstance(name, str) or not name.strip() or not isinstance(rows, list):
                raise ValueError(cls.INVALID)
            if type(table.get("enabled")) is not bool:
                raise ValueError(cls.INVALID)
            if any(not isinstance(row, dict) or keys[0] not in row for row in rows):
                raise ValueError(cls.INVALID)
            if any("regex" in row and type(row["regex"]) is not bool for row in rows):
                raise ValueError(cls.INVALID)
            normalized = {
                "name": name,
                "enabled": table["enabled"],
                "rows": [row for row in TableConfigHelper.normalize_rows(kind, rows) if row.get(keys[0])],
            }
            if kind == "prompt_dictionary":
                options = table.get("match_options", {})
                if not isinstance(options, dict) or any(type(value) is not bool for value in options.values()):
                    raise ValueError(cls.INVALID)
                normalized["match_options"] = TableConfigHelper.normalize_match_options(options)
            result.append(normalized)
        return result

    @classmethod
    def _sheet_name(cls, name, used):
        base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", name).strip(" '") or "Table"
        candidate = base[:31]
        number = 2
        while candidate.casefold() in used:
            suffix = f" ({number})"
            candidate = base[:31 - len(suffix)] + suffix
            number += 1
        used.add(candidate.casefold())
        return candidate

    @classmethod
    def _write_cell(cls, sheet, row, column, value):
        if isinstance(value, str) and len(value) > 32767:
            raise ValueError(cls.TEXT_TOO_LONG)
        cell = sheet.cell(row=row, column=column, value=value)
        if isinstance(value, str):
            # openpyxl otherwise interprets leading '=' and error-like strings.
            cell.data_type = "s"
            cell.number_format = "@"

    @classmethod
    def _save_workbook(cls, path, payload, keys, headers):
        if len(headers) != len(keys):
            raise ValueError(cls.INVALID)
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        used = {cls.METADATA_SHEET.casefold()}
        metadata_rows = []
        try:
            for table in payload["tables"]:
                sheet = workbook.create_sheet(cls._sheet_name(table["name"], used))
                sheet.freeze_panes = "A2"
                for column, header in enumerate(headers, 1):
                    cls._write_cell(sheet, 1, column, str(header))
                for index, row in enumerate(table["rows"], 2):
                    for column, key in enumerate(keys, 1):
                        cls._write_cell(sheet, index, column, row.get(key, ""))
                metadata_rows.append((
                    sheet.title, table["name"], table["enabled"],
                    json.dumps(table.get("match_options", {}), ensure_ascii=False),
                ))
            metadata = workbook.create_sheet(cls.METADATA_SHEET)
            for column, value in enumerate((cls.FORMAT, cls.VERSION, payload["kind"], json.dumps(list(keys))), 1):
                cls._write_cell(metadata, 1, column, value)
            for column, value in enumerate(cls.METADATA_HEADERS, 1):
                cls._write_cell(metadata, 2, column, value)
            for row, values in enumerate(metadata_rows, 3):
                for column, value in enumerate(values, 1):
                    cls._write_cell(metadata, row, column, value)
            metadata.sheet_state = "hidden"
            workbook.active = 0
            workbook.save(path)
        finally:
            workbook.close()

    @classmethod
    def _load_workbook(cls, path, kind, keys):
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            if cls.METADATA_SHEET not in workbook.sheetnames:
                return None
            metadata = workbook[cls.METADATA_SHEET]
            header = tuple(next(metadata.iter_rows(min_row=1, max_row=1, max_col=4, values_only=True)))
            fields = tuple(next(metadata.iter_rows(min_row=2, max_row=2, max_col=4, values_only=True)))
            if fields != cls.METADATA_HEADERS:
                raise ValueError(cls.INVALID)
            try:
                stored_keys = json.loads(header[3])
            except (TypeError, json.JSONDecodeError) as error:
                raise ValueError(cls.INVALID) from error
            payload = {"format": header[0], "version": header[1], "kind": header[2], "tables": []}
            # Validate the envelope before using its column mapping. In particular,
            # importing another category must not reinterpret its row columns.
            cls._validate_bundle({**payload, "tables": [{"name": "_", "enabled": True, "rows": []}]}, kind, keys)
            if stored_keys != list(keys):
                raise ValueError(cls.INVALID)
            seen_sheets = set()
            for values in metadata.iter_rows(min_row=3, max_col=4, values_only=True):
                if all(value is None for value in values):
                    continue
                sheet_name, name, enabled, options = values
                if not isinstance(sheet_name, str) or sheet_name in seen_sheets or sheet_name == cls.METADATA_SHEET:
                    raise ValueError(cls.INVALID)
                if sheet_name not in workbook.sheetnames:
                    raise ValueError(cls.INVALID)
                seen_sheets.add(sheet_name)
                try:
                    match_options = json.loads(options)
                except (TypeError, json.JSONDecodeError) as error:
                    raise ValueError(cls.INVALID) from error
                payload["tables"].append({
                    "name": name, "enabled": enabled, "match_options": match_options,
                    "rows": cls._read_sheet_rows(workbook[sheet_name], keys),
                })
            return cls._validate_bundle(payload, kind, keys)
        finally:
            workbook.close()

    @classmethod
    def _read_sheet_rows(cls, sheet, keys):
        rows = []
        for values in sheet.iter_rows(min_row=2, max_col=len(keys), values_only=True):
            if values[0] is None or not str(values[0]).strip():
                continue
            row = dict(zip(keys, values))
            if "regex" in row:
                value = row["regex"]
                if value is None or value == "":
                    row["regex"] = False
                elif isinstance(value, str) and value.strip().lower() in ("true", "false"):
                    row["regex"] = value.strip().lower() == "true"
                elif isinstance(value, (bool, int, float)) and value in (0, 1):
                    row["regex"] = bool(value)
                else:
                    raise ValueError(cls.INVALID)
            rows.append(row)
        return rows
