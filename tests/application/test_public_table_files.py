import json

import openpyxl
import pytest

from ModuleFolders.Application.PublicTableFileHelper import PublicTableFileHelper


TABLES = [{"name": "术语 / 保留", "enabled": False, "match_options": {"case_sensitive": True, "whole_word": False, "split_match": True}, "rows": [{"src": "=SUM(1,2)", "dst": "#REF!", "info": "中文\nnotes", "regex": False}]}]
KEYS = ("src", "dst", "info", "regex")


@pytest.mark.parametrize("extension", ["json", "xlsx"])
def test_extracted_file_helper_round_trips_legacy_multitable_bundles(tmp_path, extension):
    path = tmp_path / f"tables.{extension}"
    PublicTableFileHelper.save_bundle(path, "prompt_dictionary", TABLES, KEYS, KEYS)
    assert PublicTableFileHelper.load_bundle(path, "prompt_dictionary", KEYS) == TABLES
    if extension == "xlsx":
        workbook = openpyxl.load_workbook(path, data_only=False)
        try:
            assert workbook.active.cell(2, 1).data_type == "s"
            assert workbook.active.cell(2, 2).data_type == "s"
            assert workbook[PublicTableFileHelper.METADATA_SHEET].sheet_state == "hidden"
        finally:
            workbook.close()


@pytest.mark.parametrize("payload", [
    {"format": "ainiee.public_tables", "version": 999, "kind": "prompt_dictionary", "tables": TABLES},
    {"format": "ainiee.public_tables", "version": 1, "kind": "exclusion_list", "tables": TABLES},
    {"tables": TABLES},
])
def test_malformed_or_different_kind_bundle_is_never_interpreted_as_legacy_dictionary(tmp_path, payload):
    path = tmp_path / "bad.json"
    path.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError):
        PublicTableFileHelper.load_bundle(path, "prompt_dictionary", KEYS)


def test_failed_xlsx_export_does_not_replace_existing_file(tmp_path):
    path = tmp_path / "tables.xlsx"
    path.write_bytes(b"original")
    tables = [{**TABLES[0], "rows": [{"src": "x" * 32768}]}]
    with pytest.raises(ValueError):
        PublicTableFileHelper.save_bundle(path, "prompt_dictionary", tables, KEYS, KEYS)
    assert path.read_bytes() == b"original"
    assert list(tmp_path.glob(".tables-*")) == []
